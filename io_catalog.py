import ntpath
from collections import defaultdict, deque
from io import BytesIO
import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter
from database_manager import create_connection
from parser import (extract_io_tools, strip_lineage_annotation, split_source, location_type, resolve_path,
                    lineage_key, wildcard_match, find_paths_in_config, REVIEW_FLAGS)

# Tools that can read or write data in ways the parser cannot see. Matched as a plugin substring.
OPAQUE_PLUGINS = {
    'JupyterCode': 'Python tool',
    'RPlugin': 'R tool',
    'RunCommand': 'Run Command tool',
    'Email': 'Email tool',
    'Render': 'Render tool (writes report files)',
}
# Tools whose configuration is never a data source (containers, comments).
SKIP_SCAN_PLUGINS = ('ToolContainer', 'TextBox')

ROLE_SOURCE = 'External Source'
ROLE_INTERMEDIATE = 'Intermediate'
ROLE_FINAL = 'Final Output'

EXCEL_CELL_LIMIT = 32000


def _load_workspace(workspace_name):
    """Reads everything the catalog needs for a workspace with a single connection."""
    conn = create_connection()
    if conn is None: return None
    try:
        params = (workspace_name,)
        workspace_join = "JOIN workflows w ON t.workflow_id = w.id JOIN workspaces ws ON w.workspace_id = ws.id WHERE ws.name = ?"
        return {
            'workflows': pd.read_sql_query("""
                SELECT w.id AS workflow_db_id, w.workflow_name, w.last_parsed_at, w.source, w.source_ref,
                       w.owner, w.author, w.description, w.alteryx_version
                FROM workflows w JOIN workspaces ws ON w.workspace_id = ws.id
                WHERE ws.name = ? ORDER BY w.workflow_name
            """, conn, params=params),
            'tools': pd.read_sql_query(f"""
                SELECT w.id AS workflow_db_id, w.workflow_name, t.id AS tool_db_id, t.tool_id_xml,
                       t.plugin, t.macro, t.config_xml, t.annotation
                FROM tools t {workspace_join}
            """, conn, params=params),
            'fields': pd.read_sql_query(f"""
                SELECT tf.tool_id, tf.field_name FROM tool_fields tf
                JOIN tools t ON tf.tool_id = t.id {workspace_join} ORDER BY tf.id
            """, conn, params=params),
            'connections': pd.read_sql_query("""
                SELECT c.workflow_id, c.origin_tool_id_xml, c.destination_tool_id_xml
                FROM connections c JOIN workflows w ON c.workflow_id = w.id
                JOIN workspaces ws ON w.workspace_id = ws.id WHERE ws.name = ?
            """, conn, params=params),
        }
    finally:
        conn.close()


def _schema_finder(tools_df, fields_df, connections_df):
    """Returns a function giving the field list for a tool, walking upstream when the tool itself
    has no stored schema (Output Data tools usually don't)."""
    fields_by_tool = {tool_id: list(dict.fromkeys(group)) for tool_id, group in fields_df.groupby('tool_id')['field_name']}
    db_id_by_xml = {(row.workflow_db_id, row.tool_id_xml): row.tool_db_id for row in tools_df.itertuples()}
    tool_by_db_id = {row.tool_db_id: (row.workflow_db_id, row.tool_id_xml) for row in tools_df.itertuples()}
    origins = defaultdict(list)
    for row in connections_df.itertuples():
        origins[(row.workflow_id, row.destination_tool_id_xml)].append(row.origin_tool_id_xml)

    def find(tool_db_id):
        queue, visited = deque([tool_db_id]), set()
        while queue:
            current = queue.popleft()
            if current in visited or current not in tool_by_db_id: continue
            visited.add(current)
            if fields_by_tool.get(current): return fields_by_tool[current]
            workflow_id, xml_id = tool_by_db_id[current]
            for origin_xml in origins.get((workflow_id, xml_id), []):
                queue.append(db_id_by_xml.get((workflow_id, origin_xml)))
        return []
    return find


def _join(values):
    return '; '.join(sorted(v for v in set(values) if v))


def _role(produced, consumed):
    if produced and consumed: return ROLE_INTERMEDIATE
    if produced: return ROLE_FINAL
    return ROLE_SOURCE


@st.cache_data(show_spinner="Building I/O catalog...")
def build_io_catalog(workspace_name, aliases=()):
    """
    Builds the master I/O catalog for a workspace. Returns a dict of DataFrames:
      'I/O Detail'   - one row per input/output tool, with links to producing/consuming workflows
      'Data Stores'  - one row per unique file/table, with its producers and consumers
      'Workflows'    - one row per workflow, with its upstream and downstream workflows
      'Needs Review' - tools whose I/O could not be fully determined
    """
    data = _load_workspace(workspace_name)
    if data is None or data['tools'].empty:
        return {}
    tools_df = data['tools']
    workflows_df = data['workflows'].fillna({c: '' for c in ('source', 'source_ref', 'owner', 'author', 'description', 'alteryx_version')})
    tool_lookup = tools_df.set_index('tool_db_id').to_dict('index')
    schema_for = _schema_finder(tools_df, data['fields'], data['connections'])

    inputs, outputs = extract_io_tools(tools_df.to_dict('records'))

    # --- One row per I/O tool ---
    rows = []
    for direction, items in (('Input', inputs), ('Output', outputs)):
        for item in items:
            tool = tool_lookup[item['tool_db_id']]
            path = item.get('path') or ''
            location, obj = split_source(path)
            is_file = item['data_store_type'] == 'File'
            resolved = resolve_path(location, aliases) if is_file else location
            fields = schema_for(item['tool_db_id'])
            rows.append({
                'Workflow': item['workflow_name'],
                'Direction': direction,
                'Tool ID': tool['tool_id_xml'],
                'Tool Type': item['tool_type'],
                'Tool Annotation': strip_lineage_annotation(tool.get('annotation')),
                'Data Store Type': item['data_store_type'],
                'Location Type': location_type(location) if is_file and location else item['data_store_type'],
                'Share / Drive': ntpath.splitdrive(resolved)[0] if is_file else '',
                'Folder': ntpath.dirname(resolved) if is_file else '',
                'File Name': ntpath.basename(resolved) if is_file else '',
                'Format': ntpath.splitext(resolved)[1].lstrip('.').lower() if is_file else '',
                'Sheet / Table / Query': obj or '',
                'Resolved Path': resolved,
                'Configured Value': item.get('source_detail') or '',
                'Write Mode': item.get('write_mode') or '',
                'Flags': _join(item['flags']),
                'Field Count': len(fields),
                'Fields': ', '.join(fields),
                '_key': lineage_key(item, aliases),
                '_needs_review': bool(REVIEW_FLAGS.intersection(item['flags'])),
            })
    detail = pd.DataFrame(rows)

    # --- Link outputs to inputs across workflows ---
    producers, consumers = defaultdict(set), defaultdict(set)
    for row in rows:
        if row['_key']:
            (producers if row['Direction'] == 'Output' else consumers)[row['_key']].add(row['Workflow'])
    patterns = [k for k in consumers if '*' in k or '?' in k]
    pattern_matches = {p: [k for k in producers if k != p and wildcard_match(p, k)] for p in patterns}
    matched_by = defaultdict(list)
    for pattern, keys in pattern_matches.items():
        for key in keys: matched_by[key].append(pattern)

    def producers_of(key):
        return producers.get(key, set()).union(*(producers[k] for k in pattern_matches.get(key, [])))

    def consumers_of(key):
        return consumers.get(key, set()).union(*(consumers[p] for p in matched_by.get(key, [])))

    if not detail.empty:
        detail['Data Store Role'] = [_role(producers_of(k), consumers_of(k)) if k else '' for k in detail['_key']]
        detail['Produced By'] = [_join(producers_of(k)) if k else '' for k in detail['_key']]
        detail['Consumed By'] = [_join(consumers_of(k)) if k else '' for k in detail['_key']]
        meta = workflows_df.set_index('workflow_name')
        detail['Workflow Owner'] = detail['Workflow'].map(meta['owner']).fillna('')
        detail['Workflow Source'] = detail['Workflow'].map(meta['source']).fillna('')
        detail['Last Parsed'] = detail['Workflow'].map(meta['last_parsed_at']).fillna('')
        detail = detail.sort_values(['Workflow', 'Direction', 'Tool ID'], key=_natural_sort).reset_index(drop=True)

    # --- One row per unique data store ---
    store_rows = []
    if not detail.empty:
        for key, group in detail[detail['_key'].notna()].groupby('_key'):
            produced, consumed = producers_of(key), consumers_of(key)
            notes = []
            if len(producers.get(key, ())) > 1: notes.append('Written by more than one workflow')
            if key in pattern_matches: notes.append(f"Wildcard pattern matching {len(pattern_matches[key])} output(s) in this workspace")
            first = group.iloc[0]
            store_rows.append({
                'Data Store': group['Resolved Path'].mode().iloc[0],
                'Data Store Type': first['Data Store Type'],
                'Location Type': first['Location Type'],
                'Share / Drive': first['Share / Drive'],
                'Format': first['Format'],
                'Role': _role(produced, consumed),
                'Producer Count': len(produced),
                'Produced By': _join(produced),
                'Consumer Count': len(consumed),
                'Consumed By': _join(consumed),
                'Sheets / Tables': _join(group['Sheet / Table / Query']),
                'Notes': '; '.join(notes),
            })
    stores = pd.DataFrame(store_rows)
    if not stores.empty:
        stores = stores.sort_values(['Consumer Count', 'Data Store'], ascending=[False, True]).reset_index(drop=True)

    # --- Tools needing a human to confirm their I/O ---
    review_rows = []
    if not detail.empty:
        for row in detail[detail['_needs_review']].to_dict('records'):
            review_rows.append({'Workflow': row['Workflow'], 'Tool ID': row['Tool ID'], 'Tool Type': row['Tool Type'],
                                'Reason': row['Flags'], 'Paths Found': row['Configured Value'],
                                'Tool Annotation': row['Tool Annotation']})
    io_tool_ids = {item['tool_db_id'] for item in inputs + outputs}
    for tool in tools_df.itertuples():
        if tool.tool_db_id in io_tool_ids or any(p in tool.plugin for p in SKIP_SCAN_PLUGINS): continue
        reasons = []
        opaque = next((name for key, name in OPAQUE_PLUGINS.items() if key in tool.plugin), None)
        tool_type = opaque or tool.plugin.split('.')[-1]
        if opaque:
            reasons.append(f"{opaque} can read or write files; add a lineage annotation to catalog it")
        if tool.macro and 'Input Data Selector.yxmc' not in tool.macro:
            tool_type = f"Macro: {ntpath.basename(tool.macro)}"
            # Bare names (e.g. Cleanse.yxmc) are macros installed with Alteryx; anything with a folder is custom
            if ntpath.dirname(tool.macro):
                reasons.append('Custom macro; its contents are not parsed and may read or write data')
        paths = find_paths_in_config(tool.config_xml)
        if paths:
            reasons.append('File path referenced in configuration')
        if reasons:
            review_rows.append({'Workflow': tool.workflow_name, 'Tool ID': tool.tool_id_xml, 'Tool Type': tool_type,
                                'Reason': '; '.join(reasons), 'Paths Found': '; '.join(paths[:10]),
                                'Tool Annotation': strip_lineage_annotation(tool.annotation)})
    review = pd.DataFrame(review_rows, columns=['Workflow', 'Tool ID', 'Tool Type', 'Reason', 'Paths Found', 'Tool Annotation'])
    if not review.empty:
        review = review.sort_values(['Workflow', 'Tool ID'], key=_natural_sort).reset_index(drop=True)

    # --- One row per workflow ---
    workflow_rows = []
    for wf in workflows_df.itertuples():
        mine = detail[detail['Workflow'] == wf.workflow_name] if not detail.empty else detail
        input_keys = [k for k in mine[mine['Direction'] == 'Input']['_key'] if k] if not mine.empty else []
        output_keys = [k for k in mine[mine['Direction'] == 'Output']['_key'] if k] if not mine.empty else []
        upstream = set().union(*(producers_of(k) for k in input_keys)) - {wf.workflow_name}
        downstream = set().union(*(consumers_of(k) for k in output_keys)) - {wf.workflow_name}
        workflow_rows.append({
            'Workflow': wf.workflow_name,
            'Owner': wf.owner or '',
            'Source': wf.source or '',
            'Source Reference': wf.source_ref or '',
            'Author': wf.author or '',
            'Description': wf.description or '',
            'Alteryx Version': wf.alteryx_version or '',
            'Last Parsed': wf.last_parsed_at,
            'Inputs': int((mine['Direction'] == 'Input').sum()) if not mine.empty else 0,
            'Outputs': int((mine['Direction'] == 'Output').sum()) if not mine.empty else 0,
            'Upstream Workflows': _join(upstream),
            'Downstream Workflows': _join(downstream),
            'Items Needing Review': int((review['Workflow'] == wf.workflow_name).sum()) if not review.empty else 0,
        })

    if not detail.empty:
        detail = detail.drop(columns=['_key', '_needs_review'])
        front = ['Workflow', 'Direction', 'Tool ID', 'Tool Type', 'Tool Annotation', 'Data Store Type',
                 'Location Type', 'Share / Drive', 'Folder', 'File Name', 'Format', 'Sheet / Table / Query',
                 'Data Store Role', 'Produced By', 'Consumed By']
        detail = detail[front + [c for c in detail.columns if c not in front]]

    return {
        'I/O Detail': detail,
        'Data Stores': stores,
        'Workflows': pd.DataFrame(workflow_rows),
        'Needs Review': review,
    }


def _natural_sort(column):
    """Sorts Tool IDs numerically ('10' after '9'); other columns case-insensitively."""
    if column.name == 'Tool ID':
        return pd.to_numeric(column, errors='coerce')
    return column.astype(str).str.lower()


def to_excel_bytes(sheets):
    """Writes the catalog sheets to an .xlsx file with frozen headers, filters and sized columns."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for name, df in sheets.items():
            df = df.map(lambda v: v[:EXCEL_CELL_LIMIT] if isinstance(v, str) else v) if not df.empty else df
            df.to_excel(writer, sheet_name=name.replace('/', '-'), index=False)
            sheet = writer.sheets[name.replace('/', '-')]
            sheet.freeze_panes = 'A2'
            if not df.empty:
                sheet.auto_filter.ref = sheet.dimensions
            for idx, column in enumerate(df.columns, start=1):
                longest = max([len(str(column))] + [len(str(v)) for v in df[column].head(500)])
                sheet.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 60)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    # Keep values like '=...' (formulas in SQL or expressions) as text, not Excel formulas
                    if isinstance(cell.value, str) and cell.value.startswith('='):
                        cell.data_type = 's'
    return buffer.getvalue()
