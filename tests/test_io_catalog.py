from openpyxl import load_workbook
from io import BytesIO
import database_manager as db
import io_catalog
import reports
import tracer


def _row(df, workflow, tool_id):
    return df[(df['Workflow'] == workflow) & (df['Tool ID'] == tool_id)].iloc[0]


def test_catalog_without_mappings(workspace):
    catalog = io_catalog.build_io_catalog(workspace)
    detail = catalog['I/O Detail']
    assert len(detail) == 11

    # Excel output and input link despite different sheet syntax and case
    summary_in = _row(detail, 'Report_GL.yxmd', '2')
    assert summary_in['Produced By'] == 'Extract_GL.yxmd'
    assert summary_in['Data Store Role'] == io_catalog.ROLE_INTERMEDIATE
    assert summary_in['Sheet / Table / Query'] == 'Summary'

    # Mapped drives are not linked until a mapping is added
    assert _row(detail, 'Report_GL.yxmd', '1')['Produced By'] == ''
    gl_in = _row(detail, 'Extract_GL.yxmd', '1')
    assert gl_in['Produced By'] == ''
    assert gl_in['Fields'] == 'Account, Amount, PostDate'

    # Output Data has no schema of its own; fields come from upstream
    assert _row(detail, 'Extract_GL.yxmd', '3')['Field Count'] == 3


def test_catalog_with_drive_mapping(workspace):
    db.save_path_aliases(workspace, [('S:', r'\\fileserver\finance')])
    catalog = io_catalog.build_io_catalog(workspace, db.get_path_aliases(workspace))
    detail, workflows = catalog['I/O Detail'], catalog['Workflows']

    staging_in = _row(detail, 'Report_GL.yxmd', '1')
    assert staging_in['Produced By'] == 'Extract_GL.yxmd'
    assert staging_in['Resolved Path'] == r'\\fileserver\finance\staging\GL_Clean.yxdb'
    assert staging_in['Share / Drive'] == r'\\fileserver\finance'

    # Wildcard input (S:\raw\gl_*.csv) picks up the daily feed's dated file
    assert _row(detail, 'Extract_GL.yxmd', '1')['Produced By'] == 'Daily_Feed.yxmd'

    extract = workflows.set_index('Workflow').loc['Extract_GL.yxmd']
    assert extract['Upstream Workflows'] == 'Daily_Feed.yxmd'
    assert extract['Downstream Workflows'] == 'Report_GL.yxmd'
    assert extract['Author'] == 'Finance Analytics'
    assert extract['Source'] == 'Local Upload'


def test_data_stores_and_review(workspace):
    catalog = io_catalog.build_io_catalog(workspace)
    stores = catalog['Data Stores']
    summary = stores[stores['Data Store'].str.endswith('GL Summary.xlsx')].iloc[0]
    assert (summary['Producer Count'], summary['Consumer Count']) == (1, 1)

    review = catalog['Needs Review']
    reviewed = set(zip(review['Workflow'], review['Tool ID']))
    assert ('Report_GL.yxmd', '3') in reviewed      # dynamic input
    assert ('Report_GL.yxmd', '5') in reviewed      # relative, multi-file output
    assert ('Report_GL.yxmd', '7') in reviewed      # custom macro
    assert ('Extract_GL.yxmd', '2') in reviewed     # path in a formula
    assert ('Extract_GL.yxmd', '5') in reviewed     # Python tool without annotation
    assert ('Report_GL.yxmd', '8') not in reviewed  # built-in macro
    assert ('Report_GL.yxmd', '9') not in reviewed  # Python tool with lineage annotation
    assert ('Extract_GL.yxmd', '6') not in reviewed  # comment box


def test_excel_export(workspace):
    catalog = io_catalog.build_io_catalog(workspace)
    workbook = load_workbook(BytesIO(io_catalog.to_excel_bytes(catalog)))
    assert workbook.sheetnames == ['I-O Detail', 'Data Stores', 'Workflows', 'Needs Review']
    assert workbook['I-O Detail'].max_row == len(catalog['I/O Detail']) + 1


def test_lineage_map_uses_shared_keys(workspace):
    db.save_path_aliases(workspace, [('S:', r'\\fileserver\finance')])
    io_map = tracer.build_io_map(workspace)
    staging = io_map[r'\\fileserver\finance\staging\gl_clean.yxdb']
    assert staging['producer']['workflow_name'] == 'Extract_GL.yxmd'
    assert [c['workflow_name'] for c in staging['consumers']] == ['Report_GL.yxmd']
    assert 'not found' not in io_map

    impact = reports.generate_impact_report(reports.get_raw_io_list(workspace))
    row = impact[impact['Data Source'] == r'\\fileserver\finance\reports\gl summary.xlsx'].iloc[0]
    assert row['Producing Workflow'] == 'Extract_GL.yxmd' and row['NumberOfConsumers'] == 1


def test_upstream_trace_crosses_workflows(workspace):
    db.save_path_aliases(workspace, [('S:', r'\\fileserver\finance')])
    io_map = tracer.build_io_map(workspace)
    conn = db.create_connection()
    tool_id = conn.execute("""SELECT t.id FROM tools t JOIN workflows w ON t.workflow_id = w.id
                              WHERE w.workflow_name = 'Report_GL.yxmd' AND t.tool_id_xml = '1'""").fetchone()[0]
    workflow_id = conn.execute("SELECT id FROM workflows WHERE workflow_name = 'Report_GL.yxmd'").fetchone()[0]
    conn.close()
    lineage = tracer.trace_upstream_end_to_end(workspace, workflow_id, tool_id, 'Account', io_map)
    # Crosses both hops: the staging file from Extract_GL, then the wildcard feed from Daily_Feed
    assert {'Extract_GL.yxmd', 'Daily_Feed.yxmd'} <= set(lineage['Workflow'])
    assert lineage.iloc[0]['Workflow'] == 'Daily_Feed.yxmd'


def test_wildcard_inputs_link_in_lineage_and_impact(workspace):
    db.save_path_aliases(workspace, [('S:', r'\\fileserver\finance')])
    io_map = tracer.build_io_map(workspace)
    pattern = io_map[r'\\fileserver\finance\raw\gl_*.csv']
    assert [p['workflow_name'] for p in pattern['producers']] == ['Daily_Feed.yxmd']
    assert pattern['producer'] is None  # nothing writes the pattern itself
    dated = io_map[r'\\fileserver\finance\raw\gl_20261030.csv']
    assert [c['workflow_name'] for c in dated['consumers']] == ['Extract_GL.yxmd']

    impact = reports.generate_impact_report(reports.get_raw_io_list(workspace))
    row = impact[impact['Data Source'] == r'\\fileserver\finance\raw\gl_20261030.csv'].iloc[0]
    assert row['Producing Workflow'] == 'Daily_Feed.yxmd' and row['ConsumingWorkflows'] == ['Extract_GL.yxmd']

    conn = db.create_connection()
    workflow_id, tool_id = conn.execute("""SELECT w.id, t.id FROM tools t JOIN workflows w ON t.workflow_id = w.id
                                           WHERE w.workflow_name = 'Daily_Feed.yxmd' AND t.tool_id_xml = '1'""").fetchone()
    conn.close()
    endpoints = tracer.trace_downstream_end_to_end(workspace, workflow_id, tool_id, 'Account', io_map)
    assert 'Report_GL.yxmd' in set(endpoints['Workflow'])


def test_producers_are_matched_by_excel_sheet():
    io_map = {'book.xlsx': {'producers': [{'tool_db_id': 1, 'sheet': 'summary'}, {'tool_db_id': 2, 'sheet': 'detail'}]}}
    assert [p['tool_db_id'] for p in tracer._producers_for(io_map, 'book.xlsx', {'path': r'book.xlsx|||`Detail$`'})] == [2]
    assert len(tracer._producers_for(io_map, 'book.xlsx', {'path': 'book.xlsx'})) == 2


def test_workflow_reading_back_its_own_file_is_not_a_link(workspace):
    from parser import parse_workflow
    from conftest import FIXTURES
    import os
    tools, connections, meta = parse_workflow(os.path.join(FIXTURES, 'Diamond.yxmd'))
    db.log_workflow_details(workspace, 'Diamond.yxmd', tools, connections, meta)
    stores = io_catalog.build_io_catalog(workspace)['Data Stores']
    history = stores[stores['Data Store'].str.endswith('history.yxdb')].iloc[0]
    assert history['Role'] == io_catalog.ROLE_FINAL
    assert 'Read back by the same workflow' in history['Notes']


def test_fields_follow_the_connected_join_output(workspace):
    from parser import parse_workflow
    from conftest import FIXTURES
    import os
    tools, connections, meta = parse_workflow(os.path.join(FIXTURES, 'Join_Union.yxmd'))
    db.log_workflow_details(workspace, 'Join_Union.yxmd', tools, connections, meta)
    detail = io_catalog.build_io_catalog(workspace)['I/O Detail']
    # Output 7 is fed by the Join's R output, not its J output
    assert _row(detail, 'Join_Union.yxmd', '7')['Fields'] == 'CustomerID, Amount, OrderDate'


def test_csv_export_escapes_formulas():
    import pandas as pd
    df = pd.DataFrame({'Value': ['=HYPERLINK("x")', '+1', '@SUM(A1)', r'\\srv\a.csv'], 'Count': [1, 2, 3, 4]})
    text = io_catalog.to_csv_bytes(df).decode('utf-8-sig')
    assert "'=HYPERLINK" in text and "'+1" in text and "'@SUM" in text
    assert r'\\srv\a.csv' in text and r"'\\srv" not in text
